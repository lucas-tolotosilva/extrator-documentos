"""Exportação dos documentos para Excel, e integração mock com Google Sheets."""

from __future__ import annotations

from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Documento

router = APIRouter(prefix="/api/export", tags=["export"])

COR_CABECALHO = "1F4E78"


@router.get("/excel")
def exportar_excel(db: Session = Depends(get_db)):
    documentos = db.execute(select(Documento)).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"

    colunas = [
        "Arquivo",
        "Tipo",
        "Fornecedor",
        "CNPJ",
        "Valor",
        "Emissão",
        "Vencimento",
        "Status",
        "Confiança Média",
    ]
    ws.append(colunas)
    fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill

    for doc in documentos:
        ws.append(
            [
                doc.nome_arquivo,
                doc.tipo_documento,
                doc.fornecedor,
                doc.cnpj_fornecedor,
                doc.valor,
                doc.data_emissao.isoformat() if doc.data_emissao else None,
                doc.data_vencimento.isoformat() if doc.data_vencimento else None,
                doc.status,
                doc.confianca_media,
            ]
        )

    for i, _ in enumerate(colunas, start=1):
        ws.column_dimensions[chr(64 + i)].width = 22

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=documentos_extraidos.xlsx"},
    )


@router.post("/google-sheets")
def exportar_google_sheets():
    """Integração mock: em produção, autenticaria via OAuth2 e escreveria na planilha.

    Como este é um ambiente de demonstração sem credenciais do Google configuradas,
    retornamos uma resposta simulada indicando o modo demo.
    """

    return {
        "modo_demonstracao": True,
        "mensagem": (
            "Exportação para Google Sheets simulada (modo demonstração). "
            "Configure as credenciais OAuth2 do Google em produção para habilitar o envio real."
        ),
        "planilha_url_simulada": "https://docs.google.com/spreadsheets/d/demo-exemplo-fictício",
    }
